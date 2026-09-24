# -*- coding: utf-8 -*-
"""
Genera el libro Excel «Plan de implementación de IngenIA365ERP — COOFLOPAL».

Uso:
    python generar-plan.py                       -> escribe plan-de-implementacion-2026.xlsx junto a este script
    python generar-plan.py ruta/salida.xlsx      -> escribe en la ruta indicada
    python generar-plan.py --descargas           -> además copia el archivo a C:\\Users\\<usuario>\\Downloads

Requiere: Python 3.10+ y openpyxl. Sin macros.
Todas las fechas y reglas del plan fueron fijadas por el dueño del producto el 19/09/2026.
"""
import shutil
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

# ---------------------------------------------------------------------------
# Datos fijos del plan
# ---------------------------------------------------------------------------
HOY = date(2026, 9, 19)
COOP = "COOFLOPAL"
CONTADORA = "Contadora Rafaela Lastra España"
GESTOR = "Gestor del proyecto (por designar)"
GERENCIA = "Gerencia"
INGENIA = "Jorman Copete (IngenIA 365)"
INGENIA_DEV = "IngenIA 365 (desarrollo)"
EQ_NOM = "Equipo Nómina (2 personas, por designar)"
EQ_CON = f"{CONTADORA} + auxiliar contable (por designar)"
EQ_COM = "Equipo Comercial (2 personas, por designar)"
EQ_CAR = "Equipo Cartera (2 personas, por designar)"
ACORDADO = "Horario y modalidad: acordado entre las partes"
NO_TODO = ("No se replica toda la operación: sólo las transacciones más importantes y las que "
           "cubren todos los casos de la entidad, para interiorizar el manejo y detectar "
           "inconsistencias antes de la salida en vivo")


def D(m, d):
    return date(2026, m, d)


SABADOS = [D(9, 12) + timedelta(weeks=i) for i in range(12)]
FESTIVOS = {
    D(10, 12): "Día de la Raza",
    D(11, 2): "Todos los Santos",
    D(11, 16): "Independencia de Cartagena",
    D(12, 8): "Inmaculada Concepción",
}
CORTE = D(11, 28)
SALIDA = D(12, 1)
DIAS = {0: "lunes", 1: "martes", 2: "miércoles", 3: "jueves", 4: "viernes", 5: "sábado", 6: "domingo"}

# clave -> (nombre, relleno claro, color fuerte)
FASES = {
    "PREP": ("Preparación y parametrización", "DDEBF7", "2E75B6"),
    "CAP": ("Capacitación", "FCE4D6", "ED7D31"),
    "PAR": ("Paralelo con el aplicativo actual", "E2EFDA", "70AD47"),
    "CIERRE": ("Cierre y salida en vivo", "E4DFEC", "7030A0"),
    "GEST": ("Gestión del proyecto", "EDEDED", "7F7F7F"),
}
HITO_FILL = "FFD966"
HITO_STRONG = "BF9000"
ESTADOS = ["Pendiente", "En curso", "Terminada", "Bloqueada", "Por confirmar"]
ESTADO_FILL = {
    "Terminada": "C6EFCE",
    "En curso": "FFEB9C",
    "Bloqueada": "FFC7CE",
    "Por confirmar": "F8CBAD",
}


class T:
    """Una tarea del plan."""

    def __init__(self, id, fase, tarea, ini, fin=None, desc="", resp="", apoyo=INGENIA,
                 dep="", crit="", obs="", hito=False, estado=None):
        self.id = id
        self.fase = fase
        self.tarea = tarea
        self.ini = ini
        self.fin = fin or ini
        self.desc = desc
        self.resp = resp
        self.apoyo = apoyo
        self.dep = dep
        self.crit = crit
        self.obs = obs
        self.hito = hito
        assert self.fin >= self.ini, f"{id}: fin < inicio"
        if estado:
            self.estado = estado
        elif self.fin < HOY:
            self.estado = "Por confirmar"
        elif self.ini <= HOY <= self.fin:
            self.estado = "En curso"
        else:
            self.estado = "Pendiente"

    @property
    def fase_nombre(self):
        return FASES[self.fase][0]

    @property
    def titulo(self):
        return f"◆ HITO — {self.tarea}" if self.hito else self.tarea


def H(id, fase, tarea, fecha, dep="", obs="", resp=GESTOR):
    """Hito: tarea de un día con criterio de listo verificable por acta."""
    return T(id, fase, tarea, fecha, dep=dep, resp=resp,
             crit="Hito verificado y registrado en el acta de seguimiento", obs=obs, hito=True)


def cap(id, fecha, texto, resp, obs=ACORDADO, apoyo=INGENIA):
    return T(id, "CAP", texto, fecha, resp=resp, apoyo=apoyo,
             desc="Sesión de capacitación por proceso, con ejercicios sobre los datos reales de la cooperativa.",
             crit="Asistencia de las 2 personas del módulo; ejercicio guiado completado en el sistema", obs=obs)


# ---------------------------------------------------------------------------
# NÓMINA — sáb 12/09 → sáb 28/11 · parametrización hasta el 26/09 · nómina quincenal
# ---------------------------------------------------------------------------
OBS_S1 = "Semana ya transcurrida: confirmar el avance real con el gestor"
NOMINA = [
    T("N01", "PREP", "Acceso al sistema y recorrido inicial", D(9, 12),
      desc="Crear usuarios y segundo factor (TOTP o passkey) para las 2 personas del equipo; recorrido por el menú de Nómina y el Manual en línea (/manual).",
      resp=EQ_NOM, crit="Las 2 personas ingresan a app.ingenia365.com con usuario y segundo factor activos", obs=OBS_S1),
    T("N02", "PREP", "Plan de nómina quincenal", D(9, 14), D(9, 16),
      desc="Crear el plan en /nomina/planes con periodicidad quincenal, días de corte y fechas de pago.",
      resp=EQ_NOM, dep="N01", crit="Plan quincenal creado y aprobado por el gestor",
      obs="Confirmar si existe un segundo plan (p. ej., aprendices u honorarios mensuales). " + OBS_S1),
    T("N03", "PREP", "Parámetros legales 2026", D(9, 14), D(9, 16),
      desc="Verificar en /nomina/parametros-legales el SMMLV, el auxilio de transporte, la UVT, la tabla de retención y el FSP frente a la norma vigente.",
      resp=f"{EQ_NOM} + {CONTADORA}", dep="N01", crit="Parámetros verificados; acta breve firmada", obs=OBS_S1),
    T("N04", "PREP", "Cargos y centros de costo", D(9, 15), D(9, 17),
      desc="Crear o verificar en Maestros los cargos y centros de costo que usa la nómina actual.",
      resp=EQ_NOM, dep="N01", crit="Todos los cargos y centros de costo de la nómina actual existen en Maestros", obs=OBS_S1),
    T("N05", "PREP", "Entidades de seguridad social y parafiscales", D(9, 15), D(9, 19),
      desc="Crear EPS, ARL (con la clase de riesgo de la cooperativa), fondos de pensiones, fondos de cesantías y caja de compensación; cada entidad se vincula a su persona (tercero con NIT).",
      resp=EQ_NOM, dep="N01", crit="Todas las entidades a las que cotizan los empleados existen y tienen NIT"),
    T("N06", "PREP", "Revisión de los 40 conceptos sembrados", D(9, 17), D(9, 19),
      desc="Revisar los conceptos base (devengados, deducciones, provisiones) y desactivar los que no aplican.",
      resp=EQ_NOM, dep="N02", crit="Lista de conceptos base aprobada; los que no aplican quedan inactivos"),
    T("N07", "PREP", "Conceptos propios de la cooperativa", D(9, 21), D(9, 24),
      desc="Crear los conceptos que no vienen de fábrica: auxilios, bonificaciones, horas extra según turnos, libranzas, descuentos por aportes y créditos de la cooperativa, embargos.",
      resp=EQ_NOM, dep="N06", crit="Mapa de conceptos actual → nuevo completo: cada concepto usado en 2026 tiene su equivalente"),
    T("N08", "PREP", "Retención en la fuente por plan", D(9, 21), D(9, 23),
      desc="Configurar en /nomina/parametros-retencion el procedimiento (1 o 2) y los parámetros del plan quincenal.",
      resp=f"{EQ_NOM} + {CONTADORA}", dep="N02, N03", crit="Procedimiento configurado; caso de prueba con un empleado sujeto a retención cuadra con el actual"),
    T("N09", "PREP", "Períodos de pago", D(9, 21), D(9, 22),
      desc="Crear los períodos quincenales de septiembre a diciembre de 2026.",
      resp=EQ_NOM, dep="N02", crit="Quincenas de septiembre a diciembre creadas"),
    T("N10", "PREP", "Alta de empleados", D(9, 19), D(9, 25),
      desc="Crear la persona y la ficha de cada empleado activo desde /nomina/empleados: contrato, salario, cargo, centro de costo, EPS/AFP/ARL/CCF/cesantías, cuenta bancaria y fecha de ingreso.",
      resp=EQ_NOM, dep="N04, N05", crit="100 % de los empleados activos creados; conteo y masa salarial iguales al aplicativo actual",
      obs="Número de empleados por confirmar (dimensiona el esfuerzo)"),
    T("N11", "PREP", "Novedades fijas y recurrentes", D(9, 23), D(9, 25),
      desc="Registrar las novedades permanentes por empleado: descuentos por libranza, aportes y créditos de la cooperativa, auxilios fijos.",
      resp=EQ_NOM, dep="N07, N10", crit="Novedades permanentes registradas y verificadas contra la nómina actual"),
    T("N12", "PREP", "Verificación cruzada de la parametrización", D(9, 25), D(9, 26),
      desc="Revisar entre las 2 personas y el gestor que planes, conceptos, entidades, empleados y novedades coinciden con el aplicativo actual.",
      resp=f"{EQ_NOM} + {GESTOR}", dep="N02 a N11", crit="Lista de chequeo firmada por las 2 personas y el gestor"),
    H("N13", "PREP", "Fin de parametrización de Nómina", D(9, 26), dep="N12"),
    T("N14", "PREP", "Cuentas contables por concepto (dependencia de Contabilidad)", D(10, 19), D(10, 23),
      desc="Asignar a cada concepto activo sus cuentas contables débito/crédito con el plan de cuentas del nuevo sistema. Requisito para aprobar liquidaciones; hasta entonces sólo se calcula y compara.",
      resp=f"{EQ_NOM} + {CONTADORA}", dep="C08 (Contabilidad)",
      crit="Todos los conceptos activos tienen cuentas; una liquidación de prueba se aprueba y genera el comprobante NM",
      obs="Depende del hito C08 (plan de cuentas listo, 17/10)"),
    cap("N15", D(9, 12), "Sesión 1 (sáb 12/09): introducción, planes y conceptos", EQ_NOM, obs="Realizada: confirmar asistencia"),
    cap("N16", D(9, 19), "Sesión 2 (sáb 19/09): entidades, empleados, períodos y novedades", EQ_NOM),
    cap("N17", D(9, 26), "Sesión 3 (sáb 26/09): liquidación (calcular → revisar → aprobar), importación CSV de novedades y reportes", EQ_NOM),
    cap("N18", D(9, 30), "Sesión entre semana (mié 30/09, día por acordar): acompañamiento del ciclo de calibración", EQ_NOM),
    cap("N19", D(11, 14), "Sesión 4 (sáb 14/11): aprobación, comprobante NM y reportes para el cierre; repaso general", EQ_NOM),
    T("N20", "PAR", "Ciclo 0 (calibración): 2ª quincena de septiembre", D(9, 28), D(10, 2),
      desc="Calcular la quincena del 16 al 30/09 en el nuevo sistema sin aprobar y comparar empleado por empleado con el aplicativo actual.",
      resp=EQ_NOM, dep="N13", crit="Diferencias explicadas; parametrización ajustada",
      obs="Ciclo de prueba previo al paralelo formal de octubre y noviembre"),
    T("N21", "PAR", "Ciclo 1: novedades de la 1ª quincena de octubre", D(10, 1), D(10, 14),
      desc="Registrar en ambos sistemas las novedades del período: ausencias, horas extra, incapacidades, vacaciones.",
      resp=EQ_NOM, dep="N20", crit="Novedades del período iguales en ambos sistemas"),
    T("N22", "PAR", "Ciclo 1: liquidar la 1ª quincena de octubre", D(10, 14), D(10, 15),
      desc="Calcular y revisar la liquidación en el nuevo sistema (sin aprobar).",
      resp=EQ_NOM, dep="N21", crit="Liquidación calculada y revisada"),
    H("N23", "PAR", "Primera nómina liquidada en paralelo", D(10, 15), dep="N22", resp=EQ_NOM),
    T("N24", "PAR", "Ciclo 1: comparar con el aplicativo actual y registrar diferencias", D(10, 15), D(10, 16),
      desc="Comparar totales por concepto y por empleado; registrar cada diferencia en la bitácora con su causa y corregirla.",
      resp=EQ_NOM, dep="N22", crit="Diferencia neta = 0 o justificada en la bitácora"),
    H("N25", "PAR", "Cuadre de la primera nómina frente al aplicativo actual", D(10, 16), dep="N24", resp=EQ_NOM),
    T("N26", "PAR", "Ciclo 2: novedades de la 2ª quincena de octubre", D(10, 16), D(10, 29),
      resp=EQ_NOM, dep="N24", crit="Novedades del período iguales en ambos sistemas"),
    T("N27", "PAR", "Ciclo 2: liquidar, comparar y aprobar la 2ª quincena de octubre", D(10, 29), D(10, 31),
      desc="Calcular, comparar con el actual y, con el cuadre registrado, aprobar la liquidación: la aprobación genera el comprobante contable NM.",
      resp=EQ_NOM, dep="N14, N26", crit="Cuadre registrado; liquidación aprobada y comprobante NM generado",
      obs="Primera aprobación en el nuevo sistema (requiere N14)"),
    T("N28", "PAR", "Ciclo 2: verificar el primer comprobante NM con Contabilidad", D(11, 3), D(11, 4),
      desc="La contadora revisa el comprobante NM generado (cuentas, terceros, totales) frente a la contabilización de la nómina actual.",
      resp=f"{CONTADORA} + {EQ_NOM}", dep="N27", crit="Cuentas, terceros y totales del NM coinciden con la contabilización actual"),
    H("N29", "PAR", "Primer comprobante contable de nómina (NM) verificado", D(11, 4), dep="N28", resp=CONTADORA),
    T("N30", "PAR", "Ciclo 3: novedades de la 1ª quincena de noviembre", D(11, 3), D(11, 12),
      resp=EQ_NOM, dep="N27", crit="Novedades del período iguales en ambos sistemas", obs="El 02/11 es festivo"),
    T("N31", "PAR", "Ciclo 3: liquidar, comparar y aprobar la 1ª quincena de noviembre", D(11, 12), D(11, 14),
      resp=EQ_NOM, dep="N30", crit="Cuadre registrado; liquidación aprobada; NM generado"),
    T("N32", "PAR", "Ciclo 4: novedades de la 2ª quincena de noviembre", D(11, 17), D(11, 25),
      resp=EQ_NOM, dep="N31", crit="Novedades del período iguales en ambos sistemas", obs="El 16/11 es festivo"),
    T("N33", "PAR", "Ciclo 4: liquidar, comparar y aprobar la 2ª quincena de noviembre", D(11, 26), D(11, 28),
      resp=EQ_NOM, dep="N32", crit="Última nómina en paralelo cuadrada y aprobada; NM verificado por la contadora"),
    T("N34", "PAR", "Reportes de nómina en cada ciclo", D(10, 15), D(11, 28),
      desc="Generar en /reportes/nomina la planilla, los desprendibles y el resumen por concepto (Excel/PDF/Word) y compararlos con los del aplicativo actual.",
      resp=EQ_NOM, dep="N22", crit="Reportes de cada ciclo generados y comparados"),
    T("N35", "PAR", "Información para seguridad social (PILA)", D(10, 5), D(10, 16),
      desc="Definir con IngenIA cómo se obtiene del nuevo sistema la información para la planilla PILA.",
      resp=EQ_NOM, crit="Procedimiento documentado",
      obs="Alcance por confirmar con el proveedor; mientras tanto se sigue el procedimiento actual"),
    T("N36", "CIERRE", "Acumulados 2026 por empleado", D(11, 23), D(11, 30),
      desc="Registrar los acumulados para prestaciones (cesantías, intereses, prima, vacaciones) al 30/11 para que las provisiones y liquidaciones de diciembre salgan correctas.",
      resp=f"{EQ_NOM} + {CONTADORA}", dep="N33", crit="Acumulados al 30/11 iguales al aplicativo actual",
      obs="Confirmar con IngenIA el mecanismo de carga de acumulados"),
    T("N37", "CIERRE", "Respaldo de la nómina actual y plan de reversa", D(11, 27), D(11, 30),
      desc="Copia de seguridad de la nómina actual; exportar los informes de 2026 a PDF/Excel; escribir cómo se liquidaría la 1ª quincena de diciembre en el actual si el nuevo falla.",
      resp=f"{GESTOR} + {EQ_NOM}", crit="Copia verificada; informes exportados; procedimiento de reversa escrito"),
    H("N38", "CIERRE", "Corte del paralelo de Nómina", CORTE, dep="N33"),
    H("N39", "CIERRE", "Salida en vivo: la 1ª quincena de diciembre se liquida sólo en IngenIA365ERP", SALIDA,
      obs="Primera liquidación en vivo prevista alrededor del 15/12"),
    T("N40", "CIERRE", "Acompañamiento de la primera semana en vivo", D(12, 1), D(12, 4),
      resp=EQ_NOM, crit="Novedades de diciembre registradas sólo en el nuevo; incidencias resueltas o escaladas"),
]

# ---------------------------------------------------------------------------
# CONTABILIDAD — sáb 26/09 → sáb 28/11 · parametrización hasta el 31/10
# ---------------------------------------------------------------------------
CONTABILIDAD = [
    T("C01", "PREP", "Acceso y roles del equipo contable", D(9, 26),
      desc="Usuarios y segundo factor para la contadora y el auxiliar; roles según función (el permiso de contabilizar es aparte del de digitar).",
      resp=EQ_CON, crit="Contadora y auxiliar ingresan con MFA; sólo la contadora contabiliza"),
    T("C02", "PREP", "Decisiones de configuración", D(9, 28), D(10, 2),
      desc="Definir con la contadora: catálogo (PUC Solidario oficial de la Supersolidaria, CUIF), nivel de movimiento (5 o 6), ejercicio 2026 y sucursal principal.",
      resp=f"{CONTADORA} + {GERENCIA}", dep="C01", crit="Acta de decisiones firmada",
      obs="Decisión difícil de revertir una vez contabilizado el primer comprobante"),
    T("C03", "PREP", "Configuración inicial (cuatro ojos)", D(10, 5), D(10, 6),
      desc="Ejecutar la configuración inicial del módulo; la segunda persona aprueba (control de cuatro ojos).",
      resp=f"{CONTADORA} + {GESTOR}", dep="C02", crit="Configuración inicial ejecutada y aprobada por la segunda persona"),
    T("C04", "PREP", "Validación del catálogo PUC por la contadora", D(10, 5), D(10, 9),
      desc="Revisar el catálogo CUIF cargado (2.110 cuentas) en producción y registrar la validación en el sistema.",
      resp=CONTADORA, dep="C03", crit="Validación registrada en el sistema"),
    H("C05", "PREP", "Catálogo PUC validado por la contadora", D(10, 9), dep="C04", resp=CONTADORA),
    T("C06", "PREP", "Mapa de cuentas actual → nuevo", D(10, 5), D(10, 16),
      desc="Homologar el plan de cuentas del aplicativo actual con el CUIF e identificar las cuentas que requieren auxiliares propios.",
      resp=EQ_CON, dep="C04", crit="Hoja de homologación con el 100 % de las cuentas con saldo o movimiento en 2026"),
    T("C07", "PREP", "Plan de cuentas: auxiliares propios", D(10, 13), D(10, 17),
      desc="Crear los auxiliares (7 a 9 dígitos en nivel 5; 10 a 12 en nivel 6): bancos por cuenta, cajas, reservas, fondos sociales, provisiones, excedentes, cartera, aportes e inventarios.",
      resp=EQ_CON, dep="C06", crit="Todas las cuentas del mapa existen en el plan de cuentas", obs="El 12/10 es festivo"),
    H("C08", "PREP", "Plan de cuentas listo para Nómina, Cartera y Comercial", D(10, 17), dep="C07", resp=CONTADORA),
    T("C09", "PREP", "Tipos de comprobante", D(10, 13), D(10, 16),
      desc="Crear los tipos: ingreso/recibo de caja, egreso, causación, nota contable, nómina (NM), ajuste y apertura.",
      resp=EQ_CON, dep="C03", crit="Tipos creados con numeración definida"),
    T("C10", "PREP", "Períodos contables", D(10, 13), D(10, 16),
      desc="Abrir los períodos de 2026 necesarios (noviembre para el paralelo, diciembre para la salida en vivo) y escribir la política de cierre mensual.",
      resp=EQ_CON, dep="C03", crit="Períodos abiertos; política de cierre escrita"),
    T("C11", "PREP", "Terceros (Maestros → Personas)", D(10, 6), D(10, 23),
      desc="Verificar o crear las personas de proveedores, entidades, bancos, empleados y asociados que usan las transacciones del paralelo.",
      resp="Auxiliar contable (por designar)", dep="C01", crit="Terceros de la muestra completos con NIT/cédula y ciudad"),
    T("C12", "PREP", "Maestros transversales", D(10, 6), D(10, 9),
      desc="Ciudades, bancos, centros de costo y sucursales.",
      resp=EQ_CON, dep="C01", crit="Maestros verificados contra el aplicativo actual"),
    T("C13", "PREP", "Cuentas contables por concepto de nómina (con Nómina)", D(10, 19), D(10, 23),
      desc="Acompañar al equipo de Nómina en la asignación de cuentas por concepto (ver N14).",
      resp=f"{CONTADORA} + {EQ_NOM}", dep="C08", crit="Liquidación de prueba aprobada genera un NM correcto"),
    T("C14", "PREP", "Cuentas contables para Cartera y Comercial", D(10, 20), D(10, 30),
      desc="Definir las cuentas que usarán Cartera (créditos, intereses, mora, provisiones, aportes) e Inventario/Comercial (inventarios, costo, ventas, IVA).",
      resp=CONTADORA, dep="C08", crit="Documento de cuentas entregado a los equipos de Cartera y Comercial"),
    T("C15", "PREP", "Parametrizaciones inválidas en cero", D(10, 26), D(10, 30),
      desc="Revisar la pantalla de parametrizaciones inválidas y corregir cada pendiente.",
      resp=EQ_CON, dep="C07 a C14", crit="Pantalla sin pendientes"),
    T("C16", "PREP", "Comprobante de prueba por tipo", D(10, 27), D(10, 30),
      desc="Digitar en borradores un comprobante de cada tipo (operador) y contabilizarlo (contadora) con validación por campo; anular según el protocolo.",
      resp=EQ_CON, dep="C09, C15", crit="Un comprobante por tipo contabilizado y anulado; protocolo de cuatro ojos probado"),
    H("C17", "PREP", "Fin de parametrización de Contabilidad", D(10, 31), dep="C15, C16", resp=CONTADORA),
    cap("C18", D(9, 26), "Sesión 1 (sáb 26/09): introducción — catálogo CUIF, niveles, cuatro ojos y plan del módulo", EQ_CON),
    cap("C19", D(10, 3), "Sesión 2 (sáb 03/10): configuración inicial, catálogos y validación del catálogo", EQ_CON),
    cap("C20", D(10, 10), "Sesión 3 (sáb 10/10): plan de cuentas, tipos de comprobante y períodos", EQ_CON),
    cap("C21", D(10, 17), "Sesión 4 (sáb 17/10): comprobantes y borradores — causaciones, pagos y egresos con casos de la cooperativa", EQ_CON),
    cap("C22", D(10, 24), "Sesión 5 (sáb 24/10): nómina (NM), cuentas de cartera e inventario; protocolo del paralelo", EQ_CON),
    cap("C23", D(10, 31), "Sesión 6 (sáb 31/10): cierre de parametrización, muestra del paralelo y bitácora de diferencias", EQ_CON),
    cap("C24", D(11, 11), "Sesión entre semana (mié 11/11, día por acordar): revisión de los comprobantes del ciclo 1 con la contadora", EQ_CON),
    T("C25", "PAR", "Muestra de transacciones del paralelo", D(11, 3), D(11, 4),
      desc="Seleccionar lo que se replica: causaciones de compras y servicios, egresos y pagos, recibos de caja, movimientos bancarios, notas de ajuste, nómina (NM) y un caso de cada situación especial (fondos sociales, excedentes, provisiones).",
      resp=f"{CONTADORA} + {GESTOR}", dep="C17", crit="Lista con al menos una transacción por tipo de comprobante y por cuenta principal", obs=NO_TODO),
    T("C26", "PAR", "Ciclo 1: digitar y contabilizar la muestra (1 al 14/11)", D(11, 3), D(11, 13),
      desc="El operador digita en borradores; la contadora revisa y contabiliza.",
      resp=EQ_CON, dep="C25", crit="Comprobantes de la muestra contabilizados"),
    T("C27", "PAR", "Ciclo 1: comparar con el aplicativo actual y registrar diferencias", D(11, 13), D(11, 14),
      desc="Comparar saldos y movimientos de las cuentas de la muestra; registrar cada diferencia en la bitácora con causa y acción.",
      resp=EQ_CON, dep="C26", crit="Diferencias registradas y con responsable"),
    T("C28", "PAR", "Verificar los comprobantes NM de noviembre", D(11, 17), D(11, 28),
      desc="Revisar los NM de las dos quincenas de noviembre generados por Nómina.",
      resp=CONTADORA, dep="N31, N33 (Nómina)", crit="NM revisados y cuadrados con la contabilización actual"),
    T("C29", "PAR", "Ciclo 2: digitar y contabilizar la muestra (15 al 27/11)", D(11, 17), D(11, 27),
      resp=EQ_CON, dep="C27", crit="Comprobantes de la muestra contabilizados", obs="El 16/11 es festivo"),
    T("C30", "PAR", "Ciclo 2: comparar y registrar diferencias", D(11, 27), D(11, 28),
      resp=EQ_CON, dep="C29", crit="Diferencias registradas; sin diferencias abiertas sin plan"),
    T("C31", "PAR", "Ajustes de parametrización derivados del paralelo", D(11, 14), D(11, 27),
      resp=EQ_CON, dep="C27", crit="Cada diferencia de la bitácora cerrada con ajuste o justificación"),
    T("C32", "PAR", "Entrega E2: informes contables y estados financieros (dependencia del proveedor)", D(11, 2), D(11, 28),
      desc="IngenIA entrega balance de prueba, libros (diario, mayor, auxiliares), ESF y ERI. Fecha por confirmar por el proveedor (ver G07).",
      resp=INGENIA_DEV, apoyo=INGENIA_DEV, crit="Entrega instalada en producción",
      obs="Sin E2, la comparación de saldos del paralelo se hace consultando comprobantes y exportando"),
    T("C33", "PAR", "Comparar el balance de prueba del paralelo con el actual (requiere E2)", D(11, 23), D(11, 28),
      resp=CONTADORA, dep="C32", crit="Balance de prueba de la muestra igual en ambos sistemas", obs="Depende de la fecha de E2"),
    T("C34", "PAR", "Entrega E3: integración contable automática de Cartera, Inventario/Comercial, Tesorería y CDT (dependencia del proveedor)", D(11, 2), D(11, 28),
      desc="Fecha por confirmar por IngenIA (ver G07). Mientras no exista, los movimientos de esos módulos se contabilizan manualmente con comprobantes resumen.",
      resp=INGENIA_DEV, apoyo=INGENIA_DEV, crit="Fecha comprometida registrada en este plan; decisión sobre el procedimiento manual documentada"),
    T("C35", "PAR", "Procedimiento de contabilización manual de Cartera y Comercial", D(11, 9), D(11, 20),
      desc="Definir y probar el comprobante resumen (diario o mensual) para ventas, inventario, cartera y aportes hasta que llegue E3.",
      resp=CONTADORA, dep="C14, C34", crit="Comprobante resumen definido, probado y documentado"),
    T("C36", "CIERRE", "Balance de prueba del aplicativo actual al 30/11", D(11, 30), D(12, 2),
      desc="Obtener el balance por auxiliar al 30/11 con detalle por tercero en bancos, CxC, CxP, cartera y aportes.",
      resp=CONTADORA, crit="Balance impreso/exportado y firmado"),
    T("C37", "CIERRE", "Comprobante de saldos iniciales al 30/11", D(12, 1), D(12, 4),
      desc="Digitar el comprobante de apertura en el nuevo sistema usando el mapa de cuentas; contabilizar con cuatro ojos.",
      resp=EQ_CON, dep="C36, C07", crit="Balance de prueba del nuevo sistema = balance del actual al 30/11",
      obs="Si el actual cierra noviembre después del 01/12, se registra un comprobante de ajuste (fecha por confirmar)"),
    H("C38", "CIERRE", "Saldos al 30/11 cargados y conciliados", D(12, 4), dep="C37, C39", resp=CONTADORA),
    T("C39", "CIERRE", "Conciliación de saldos por módulo", D(12, 2), D(12, 4),
      desc="Cruzar cartera, aportes, inventario, acumulados de nómina y bancos con los saldos cargados en cada módulo.",
      resp=CONTADORA, dep="C37", crit="Conciliación firmada por módulo"),
    T("C40", "CIERRE", "Respaldo del aplicativo actual y plan de reversa", D(11, 27), D(11, 30),
      desc="Copia completa de la base de datos e informes 2026 en PDF; el aplicativo actual se conserva en solo lectura; criterio escrito para volver atrás.",
      resp=f"{GESTOR} + {INGENIA}", crit="Copia verificada; criterio de reversa aprobado por gerencia"),
    H("C41", "CIERRE", "Corte del paralelo", CORTE, dep="C30"),
    H("C42", "CIERRE", "Salida en vivo", SALIDA),
    T("C43", "CIERRE", "Acompañamiento de la primera semana en vivo", D(12, 1), D(12, 4),
      resp=EQ_CON, crit="Comprobantes de diciembre sólo en el nuevo; incidencias resueltas o escaladas"),
]

# ---------------------------------------------------------------------------
# COMERCIAL (menú Inventario + Tesorería) — sáb 26/09 → sáb 28/11 · parametrización hasta el 31/10
# ---------------------------------------------------------------------------
COMERCIAL = [
    T("M01", "PREP", "Acceso y roles del equipo comercial", D(9, 26),
      resp=EQ_COM, crit="Las 2 personas ingresan con MFA; roles asignados"),
    T("M02", "PREP", "Levantamiento de la operación comercial", D(9, 28), D(10, 2),
      desc="Documentar puntos de venta, bodegas, turnos, vendedores, listas de precios, tipos de movimiento y volumen de productos del aplicativo actual.",
      resp=f"{EQ_COM} + {GESTOR}", dep="M01", crit="Documento de levantamiento aprobado por el gestor"),
    T("M03", "PREP", "Grupos primarios y grupos de productos", D(10, 5), D(10, 7), resp=EQ_COM, dep="M02", crit="Estructura de grupos igual o mejorada frente a la actual"),
    T("M04", "PREP", "Bodegas y ubicaciones", D(10, 5), D(10, 7), resp=EQ_COM, dep="M02", crit="Todas las bodegas físicas creadas con sus ubicaciones"),
    T("M05", "PREP", "Puntos de venta y turnos", D(10, 6), D(10, 9), resp=EQ_COM, dep="M02", crit="Puntos de venta y esquema de turnos parametrizados"),
    T("M06", "PREP", "Vendedores y comisiones", D(10, 8), D(10, 14),
      resp=EQ_COM, dep="M02", crit="Vendedores vinculados a personas; esquema de comisiones parametrizado", obs="El 12/10 es festivo"),
    T("M07", "PREP", "Tipos de descuento y listas de precios", D(10, 13), D(10, 17), resp=EQ_COM, dep="M02", crit="Listas de precios vigentes creadas; descuentos autorizados definidos"),
    T("M08", "PREP", "Tipos de movimiento", D(10, 13), D(10, 15),
      desc="Entradas por compra, salidas por venta, traslados, ajustes, devoluciones y mermas.",
      resp=EQ_COM, dep="M02", crit="Un tipo por cada movimiento que hoy se registra"),
    T("M09", "PREP", "Catálogo de productos", D(10, 13), D(10, 23),
      desc="Crear los productos activos: código, grupo, unidad, IVA, costo, precios por lista y bodega.",
      resp=EQ_COM, dep="M03, M07", crit="100 % de los productos activos creados; conteo igual al actual",
      obs="Volumen por confirmar; si supera ~300 referencias, evaluar importación con IngenIA"),
    T("M10", "PREP", "Tesorería: conceptos, bancos y cheques", D(10, 19), D(10, 23), resp=EQ_COM, dep="M01", crit="Conceptos de tesorería, bancos y chequeras configurados"),
    T("M11", "PREP", "Terceros: clientes y proveedores (con Contabilidad)", D(10, 19), D(10, 23),
      resp=EQ_COM, dep="C11 (Contabilidad)", crit="Clientes a crédito y proveedores frecuentes creados como personas"),
    T("M12", "PREP", "Inventario inicial de prueba al 31/10", D(10, 29), D(10, 31),
      desc="Cargar las existencias por bodega al 31/10 (conteo o saldo del actual) para las referencias de la muestra, como base del paralelo.",
      resp=EQ_COM, dep="M09", crit="Inventario valorizado de la muestra comparable con el actual"),
    T("M13", "PREP", "Verificación cruzada de la parametrización", D(10, 27), D(10, 30),
      resp=f"{EQ_COM} + {GESTOR}", dep="M03 a M11", crit="Lista de chequeo firmada"),
    H("M14", "PREP", "Fin de parametrización de Comercial", D(10, 31), dep="M12, M13"),
    cap("M15", D(9, 26), "Sesión 1 (sáb 26/09): introducción al módulo Inventario/Comercial y Tesorería; plan del módulo", EQ_COM),
    cap("M16", D(10, 3), "Sesión 2 (sáb 03/10): grupos, bodegas, ubicaciones, puntos de venta y turnos", EQ_COM),
    cap("M17", D(10, 10), "Sesión 3 (sáb 10/10): productos, listas de precios, descuentos, vendedores y comisiones", EQ_COM),
    cap("M18", D(10, 17), "Sesión 4 (sáb 17/10): tipos de movimiento, movimientos, kardex e inventario valorizado", EQ_COM),
    cap("M19", D(10, 24), "Sesión 5 (sáb 24/10): facturación y tesorería (cheques, facturas CxP/CxC, flujo de caja)", EQ_COM),
    cap("M20", D(10, 31), "Sesión 6 (sáb 31/10): cierre de parametrización, muestra del paralelo y bitácora", EQ_COM),
    cap("M21", D(11, 12), "Sesión entre semana (jue 12/11, día por acordar): acompañamiento en punto de venta durante el ciclo 1", EQ_COM),
    T("M22", "PAR", "Muestra de transacciones del paralelo", D(11, 3), D(11, 4),
      desc="Ventas por punto de venta y turno (contado y crédito), compras, traslados entre bodegas, devoluciones, ajustes, facturas CxP/CxC y cheques.",
      resp=f"{EQ_COM} + {GESTOR}", dep="M14", crit="Al menos un caso por tipo de movimiento y por punto de venta", obs=NO_TODO),
    T("M23", "PAR", "Ciclo 1: registrar los movimientos de la muestra (1 al 14/11)", D(11, 3), D(11, 13), resp=EQ_COM, dep="M22", crit="Movimientos de la muestra registrados"),
    T("M24", "PAR", "Ciclo 1: comparar kardex, inventario valorizado y ventas con el actual; registrar diferencias", D(11, 13), D(11, 14),
      resp=EQ_COM, dep="M23", crit="Diferencias registradas en la bitácora con causa y acción"),
    T("M25", "PAR", "Ciclo 2: registrar los movimientos de la muestra (15 al 27/11)", D(11, 17), D(11, 27), resp=EQ_COM, dep="M24", crit="Movimientos de la muestra registrados", obs="El 16/11 es festivo"),
    T("M26", "PAR", "Ciclo 2: comparar y registrar diferencias", D(11, 27), D(11, 28), resp=EQ_COM, dep="M25", crit="Sin diferencias abiertas sin plan"),
    T("M27", "PAR", "Facturación: datos del emisor, numeración e impuestos", D(11, 3), D(11, 13),
      desc="Emitir facturas de prueba y verificar datos del emisor, numeración e IVA.",
      resp=EQ_COM, dep="M14", crit="Facturas de prueba correctas; alcance de facturación electrónica (DIAN) confirmado con el proveedor",
      obs="Punto por confirmar con IngenIA"),
    T("M28", "PAR", "Cierres de turno y flujo de caja de la muestra", D(11, 10), D(11, 27), resp=EQ_COM, dep="M23", crit="Cierres de turno cuadrados con el actual"),
    T("M29", "PAR", "Ajustes de parametrización derivados del paralelo", D(11, 14), D(11, 27), resp=EQ_COM, dep="M24", crit="Cada diferencia cerrada con ajuste o justificación"),
    T("M30", "PAR", "Contabilización de ventas e inventario (comprobante resumen hasta E3)", D(11, 9), D(11, 20),
      resp=f"{EQ_COM} + {CONTADORA}", dep="C35 (Contabilidad)", crit="Comprobante resumen de la muestra contabilizado"),
    T("M31", "CIERRE", "Toma física de inventario al 30/11", D(11, 28), D(11, 30),
      desc="Conteo por bodega al cierre del 30/11 (o del 28/11 con ajuste de los movimientos del 29 y 30).",
      resp=EQ_COM, crit="Conteo por bodega firmado; diferencias con el kardex del actual explicadas"),
    T("M32", "CIERRE", "Carga del inventario inicial definitivo", D(11, 30), D(12, 2),
      resp=EQ_COM, dep="M31", crit="Inventario valorizado = toma física; diferencia con contabilidad conciliada con la contadora"),
    T("M33", "CIERRE", "Saldos de CxC clientes y CxP proveedores al 30/11", D(12, 1), D(12, 3),
      resp=EQ_COM, dep="M11", crit="Facturas pendientes cargadas en Tesorería; totales iguales al actual"),
    T("M34", "CIERRE", "Respaldo y plan de reversa (comercial)", D(11, 27), D(11, 30),
      resp=f"{GESTOR} + {EQ_COM}", crit="Informes de inventario y ventas 2026 exportados; procedimiento de reversa escrito"),
    H("M35", "CIERRE", "Corte del paralelo", CORTE, dep="M26"),
    H("M36", "CIERRE", "Salida en vivo: ventas y movimientos sólo en el nuevo", SALIDA),
    T("M37", "CIERRE", "Acompañamiento de la primera semana en vivo", D(12, 1), D(12, 4),
      resp=EQ_COM, crit="Ventas y movimientos de diciembre sólo en el nuevo; incidencias resueltas o escaladas"),
]

# ---------------------------------------------------------------------------
# CARTERA FINANCIERA — sáb 24/10 → sáb 28/11 · parametrización hasta el mar 10/11 · créditos y aportes
# ---------------------------------------------------------------------------
CARTERA = [
    T("K01", "PREP", "Acceso y roles del equipo de cartera", D(10, 24),
      resp=EQ_CAR, crit="Las 2 personas ingresan con MFA; roles asignados"),
    T("K02", "PREP", "Levantamiento de productos de crédito y aportes", D(10, 26), D(10, 28),
      desc="Documentar las líneas vigentes (tasas, plazos, garantías, periodicidad), el reglamento de aportes, la política de provisión y calificación, zonas y agencias.",
      resp=f"{EQ_CAR} + {GERENCIA}", dep="K01", crit="Documento de levantamiento aprobado"),
    T("K03", "PREP", "Agencias, zonas y tipos de zona", D(10, 27), D(10, 28), resp=EQ_CAR, dep="K02", crit="Agencias y zonas creadas"),
    T("K04", "PREP", "Periodicidad y códigos de movimiento", D(10, 27), D(10, 29), resp=EQ_CAR, dep="K02", crit="Periodicidades y códigos de movimiento equivalentes a los actuales"),
    T("K05", "PREP", "Tasas de interés y tasas por plazo", D(10, 29), D(10, 30),
      resp=EQ_CAR, dep="K02", crit="Tasas por línea y plazo y tasa de mora dentro del límite de usura vigente"),
    T("K06", "PREP", "Líneas de crédito", D(10, 29), D(11, 3), resp=EQ_CAR, dep="K04, K05", crit="Todas las líneas vigentes creadas", obs="El 02/11 es festivo"),
    T("K07", "PREP", "Parámetros de provisión y calificación", D(11, 3), D(11, 4),
      resp=f"{EQ_CAR} + {CONTADORA}", dep="K02",
      crit="Categorías A–E por días de mora y porcentajes de provisión según la Circular Básica Contable y Financiera de la Supersolidaria"),
    T("K08", "PREP", "Parámetros de aportes sociales", D(11, 3), D(11, 4), resp=EQ_CAR, dep="K02", crit="Cuota, periodicidad y forma de recaudo de aportes configuradas"),
    T("K09", "PREP", "Conceptos de descuento y estados de retiro", D(11, 4), D(11, 5), resp=EQ_CAR, dep="K02", crit="Conceptos y estados equivalentes a los actuales"),
    T("K10", "PREP", "SIPLA / SARLAFT: parámetros", D(11, 5), D(11, 6),
      resp=f"{EQ_CAR} + oficial de cumplimiento (si existe)", dep="K02", crit="Umbrales y señales según el manual de la cooperativa"),
    T("K11", "PREP", "Scoring y parámetros de vivienda (si aplican)", D(11, 5), D(11, 6),
      resp=EQ_CAR, dep="K02", crit="Configurados o marcados como «no aplica»", obs="Marcar «no aplica» si la cooperativa no los usa"),
    T("K12", "PREP", "Cuentas de cartera (con Contabilidad)", D(11, 5), D(11, 9),
      desc="Asignar las cuentas contables por línea y concepto: capital, intereses, mora, provisión y aportes.",
      resp=f"{EQ_CAR} + {CONTADORA}", dep="C14 (Contabilidad)", crit="Todas las líneas y conceptos con cuentas asignadas"),
    T("K13", "PREP", "Registro de asociados para la muestra", D(11, 3), D(11, 10),
      desc="Crear las personas y el registro de asociados (/asociados/registro) de la muestra del paralelo.",
      resp=EQ_CAR, dep="K01", crit="Asociados de la muestra creados", obs="Volumen supuesto < 500 asociados (por confirmar)"),
    T("K14", "PREP", "Descuento por nómina: configuración", D(11, 9), D(11, 10),
      resp=EQ_CAR, dep="K08, N14 (Nómina)", crit="Descuentos de cuotas y aportes de empleados asociados vinculados con Nómina"),
    T("K15", "PREP", "Verificación cruzada de la parametrización", D(11, 9), D(11, 10),
      resp=f"{EQ_CAR} + {GESTOR}", dep="K03 a K14", crit="Lista de chequeo firmada"),
    H("K16", "PREP", "Fin de parametrización de Cartera", D(11, 10), dep="K15", obs="Martes 10/11"),
    cap("K17", D(10, 24), "Sesión 1 (sáb 24/10): introducción y parámetros (líneas, tasas, códigos, periodicidad, zonas)", EQ_CAR),
    cap("K18", D(10, 31), "Sesión 2 (sáb 31/10): asociados, aportes, solicitudes de crédito y aprobación", EQ_CAR),
    cap("K19", D(11, 7), "Sesión 3 (sáb 07/11): desembolsos, proyecciones y plan de pagos, recaudos, descuento por nómina", EQ_CAR),
    cap("K20", D(11, 14), "Sesión 4 (sáb 14/11): mora y cobro, causación de intereses, calificación de cartera, extractos, cartera por edades e informes", EQ_CAR),
    cap("K21", D(11, 21), "Sesión 5 (sáb 21/11): revisión del paralelo y preparación de la carga de saldos", EQ_CAR),
    cap("K22", D(11, 4), "Sesión entre semana (mié 04/11, día por acordar): acompañamiento en parametrización", EQ_CAR),
    T("K23", "PAR", "Muestra de operaciones del paralelo", D(11, 11), D(11, 12),
      desc="Solicitud → aprobación → desembolso de un crédito por línea; recaudos de cuotas (caja y descuento por nómina); abonos extraordinarios; recaudo de aportes; un crédito en mora; un retiro de asociado.",
      resp=f"{EQ_CAR} + {GESTOR}", dep="K16", crit="Lista de la muestra aprobada", obs=NO_TODO),
    T("K24", "PAR", "Ciclo 1: registrar las operaciones de la muestra (11 al 20/11)", D(11, 11), D(11, 20), resp=EQ_CAR, dep="K23", crit="Operaciones registradas", obs="El 16/11 es festivo"),
    T("K25", "PAR", "Ciclo 1: comparar con el actual y registrar diferencias", D(11, 20), D(11, 21),
      resp=EQ_CAR, dep="K24", crit="Saldos, cuotas, intereses y proyecciones iguales o con diferencia justificada"),
    T("K26", "PAR", "Ciclo 2: registrar las operaciones de la muestra (23 al 27/11)", D(11, 23), D(11, 27), resp=EQ_CAR, dep="K25", crit="Operaciones registradas"),
    T("K27", "PAR", "Ciclo 2: causación de intereses y calificación de cartera al corte", D(11, 26), D(11, 27),
      resp=EQ_CAR, dep="K26", crit="Causación y calificación de la muestra ejecutadas y comparadas con el actual"),
    T("K28", "PAR", "Ciclo 2: comparar y registrar diferencias; cartera por edades y extractos", D(11, 27), D(11, 28),
      resp=EQ_CAR, dep="K27", crit="Cartera por edades y extractos de la muestra iguales al actual"),
    T("K29", "PAR", "Ajustes de parametrización derivados del paralelo", D(11, 21), D(11, 27), resp=EQ_CAR, dep="K25", crit="Cada diferencia cerrada con ajuste o justificación"),
    T("K30", "PAR", "Contabilización de cartera y aportes (comprobante resumen hasta E3)", D(11, 17), D(11, 27),
      resp=f"{EQ_CAR} + {CONTADORA}", dep="C35 (Contabilidad)", crit="Comprobante resumen de la muestra contabilizado"),
    T("K31", "PAR", "Completar el registro del 100 % de asociados", D(11, 11), D(11, 27),
      resp=EQ_CAR, dep="K13", crit="Todos los asociados activos con persona y registro creados; conteo igual al actual"),
    T("K32", "CIERRE", "Listado de créditos vigentes y aportes al 27/11 desde el actual", D(11, 25), D(11, 27),
      desc="Por asociado: saldo de capital, intereses causados, cuotas pendientes, días de mora, calificación y saldo de aportes.",
      resp=EQ_CAR, crit="Listado exportado y firmado"),
    T("K33", "CIERRE", "Precarga de créditos vigentes y aportes con saldos al 27/11", D(11, 28), D(11, 30),
      desc="Digitar en el nuevo sistema los créditos vigentes y los aportes de cada asociado (2 personas).",
      resp=EQ_CAR, dep="K31, K32", crit="Créditos y aportes digitados (supuesto < 300 créditos)",
      obs="Si el volumen es mayor, importación asistida por IngenIA"),
    T("K34", "CIERRE", "Ajuste con los movimientos del 28 al 30/11 y conciliación al 30/11", D(12, 1), D(12, 2),
      desc="Registrar los pagos y desembolsos del 28 al 30/11; cruzar la cartera por edades y el total de aportes con el actual y con el saldo contable.",
      resp=f"{EQ_CAR} + {CONTADORA}", dep="K33", crit="Cartera por edades y total de aportes iguales al actual y al saldo contable al 30/11"),
    H("K35", "CIERRE", "Saldos de cartera y aportes al 30/11 conciliados", D(12, 2), dep="K34", resp=CONTADORA),
    T("K36", "CIERRE", "Respaldo y plan de reversa (cartera)", D(11, 27), D(11, 30),
      resp=f"{GESTOR} + {EQ_CAR}", crit="Informes de cartera y aportes 2026 exportados; procedimiento de reversa escrito"),
    H("K37", "CIERRE", "Corte del paralelo", CORTE, dep="K28"),
    H("K38", "CIERRE", "Salida en vivo", SALIDA),
    T("K39", "CIERRE", "Acompañamiento de la primera semana en vivo", D(12, 1), D(12, 4),
      resp=EQ_CAR, crit="Recaudos y desembolsos de diciembre sólo en el nuevo; incidencias resueltas o escaladas"),
]

# ---------------------------------------------------------------------------
# GESTIÓN DEL PROYECTO (transversal)
# ---------------------------------------------------------------------------
GESTION = [
    T("G01", "GEST", "Designar gestor del proyecto y equipos (mínimo 2 personas por módulo)", D(9, 19), D(9, 23),
      resp=GERENCIA, crit="Acta con nombres, suplentes y dedicación; este libro actualizado con los nombres"),
    T("G02", "GEST", "Acordar horario, modalidad y calendario de capacitaciones", D(9, 19),
      resp=f"{GERENCIA} + {INGENIA}", crit="Acuerdo escrito: sábados por proceso, sesiones entre semana y herramienta virtual",
      obs="Acordado entre las partes (19/09)"),
    T("G03", "GEST", "Usuarios, roles y segundo factor para todos los designados", D(9, 21), D(9, 25),
      resp="Administrador de la cooperativa", dep="G01", crit="Cada persona ingresa con MFA; roles según función (el Operador digita, no contabiliza)"),
    T("G04", "GEST", "Comunicado interno de inicio del proyecto", D(9, 22), D(9, 26),
      resp=GERENCIA, dep="G01", crit="Todo el personal conoce el alcance, las fechas, los equipos y qué cambia el 01/12"),
    T("G05", "GEST", "Bitácora central de inconsistencias y decisiones", D(9, 28), CORTE,
      resp=GESTOR, crit="Cada diferencia del paralelo con fecha, módulo, causa, responsable y cierre"),
    T("G06", "GEST", "Seguimiento semanal del plan", D(9, 26), CORTE,
      desc="Reunión corta cada sábado (tras la capacitación) para actualizar el estado de las tareas de este libro y escalar desvíos a gerencia.",
      resp=f"{GESTOR} + {INGENIA}", crit="Estado actualizado cada semana; desvíos escalados"),
    T("G07", "GEST", "Confirmar las fechas de las entregas E2 y E3 con el proveedor", D(9, 28), D(10, 2),
      resp=GESTOR, crit="Fechas registradas en el plan (C32, C34) y en el acta de seguimiento"),
    T("G08", "GEST", "Comunicado de salida en vivo al personal", D(11, 23), D(11, 27),
      resp=GERENCIA, crit="Todo el personal sabe que desde el 01/12 sólo se usa IngenIA365ERP"),
    T("G09", "GEST", "Comité de decisión de salida en vivo (go / no-go)", CORTE,
      resp=f"{GERENCIA} + {GESTOR} + {INGENIA}", dep="N38, C41, M35, K37",
      crit="Acta con decisión. Criterios: paralelo cuadrado en los 4 módulos, saldos listos, respaldo hecho, equipos capacitados"),
    H("G10", "GEST", "Salida en vivo — sólo IngenIA365ERP", SALIDA, dep="G09", resp=GERENCIA),
    T("G11", "GEST", "Soporte intensivo y revisión posterior a la salida", D(12, 1), D(12, 4),
      resp=f"{GESTOR} + {INGENIA}", crit="Incidencias de la primera semana resueltas o con plan; lecciones aprendidas registradas"),
]

MODULOS = [
    # (nombre hoja, tareas, inicio, fin, fin parametrización)
    ("Nómina", NOMINA, D(9, 12), CORTE, D(9, 26)),
    ("Contabilidad", CONTABILIDAD, D(9, 26), CORTE, D(10, 31)),
    ("Comercial", COMERCIAL, D(9, 26), CORTE, D(10, 31)),
    ("Cartera financiera", CARTERA, D(10, 24), CORTE, D(11, 10)),
]

# ---------------------------------------------------------------------------
# Estilos
# ---------------------------------------------------------------------------
NAVY = "1F3864"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F_TITLE = Font(name="Calibri", size=16, bold=True, color=NAVY)
F_SUB = Font(name="Calibri", size=11, italic=True, color="595959")
F_H2 = Font(name="Calibri", size=13, bold=True, color=NAVY)
F_HEAD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
F_BOLD = Font(name="Calibri", size=11, bold=True)
F_NORMAL = Font(name="Calibri", size=11)
F_WARN = Font(name="Calibri", size=12, bold=True, color="C00000")
FILL_HEAD = PatternFill("solid", start_color=NAVY, end_color=NAVY)
FILL_WARN = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
A_WRAP = Alignment(wrap_text=True, vertical="top")
A_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATE_FMT = "dd/mm/yyyy"


def fill(hex_color):
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_header(ws, row, headers, start_col=1):
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + j, value=h)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = A_CENTER
        c.border = BORDER
    ws.row_dimensions[row].height = 32


def write_legend(ws, row, start_col=1):
    """Fila de leyenda de fases con colores."""
    c = ws.cell(row=row, column=start_col, value="Leyenda de fases:")
    c.font = F_BOLD
    c.alignment = Alignment(horizontal="right", vertical="center")
    col = start_col + 1
    for key, (nombre, light, strong) in FASES.items():
        c = ws.cell(row=row, column=col, value=nombre)
        c.fill = fill(light)
        c.font = Font(name="Calibri", size=10, bold=True, color=strong)
        c.alignment = A_CENTER
        c.border = BORDER
        col += 1
    c = ws.cell(row=row, column=col, value="◆ Hito")
    c.fill = fill(HITO_FILL)
    c.font = Font(name="Calibri", size=10, bold=True, color=HITO_STRONG)
    c.alignment = A_CENTER
    c.border = BORDER
    ws.row_dimensions[row].height = 30


def add_estado_validation(ws, col_letter, first_row, last_row):
    dv = DataValidation(type="list", formula1='"' + ",".join(ESTADOS) + '"', allow_blank=True)
    dv.error = "Elija un estado de la lista"
    dv.errorTitle = "Estado"
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{first_row}:{col_letter}{last_row}")
    for estado, color in ESTADO_FILL.items():
        ws.conditional_formatting.add(
            f"{col_letter}{first_row}:{col_letter}{last_row}",
            CellIsRule(operator="equal", formula=[f'"{estado}"'], fill=fill(color)),
        )


def print_setup(ws, title_rows):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if title_rows:
        ws.print_title_rows = title_rows


# ---------------------------------------------------------------------------
# Hoja por módulo
# ---------------------------------------------------------------------------
COLS_MOD = ["Nº", "Fase", "Tarea", "Descripción", "Responsable (cooperativa)", "Apoyo IngenIA",
            "Fecha inicio", "Fecha fin", "Duración (días)", "Dependencia",
            "Entregable / criterio de listo", "Estado", "Observaciones"]


def build_modulo(wb, nombre, tareas, ini, fin, fin_param):
    ws = wb.create_sheet(nombre)
    ws.sheet_properties.tabColor = FASES["PREP"][2]
    ws["A1"] = f"{nombre} — Plan de implementación IngenIA365ERP · {COOP}"
    ws["A1"].font = F_TITLE
    ws["A2"] = (f"Del {ini:%d/%m/%Y} al {fin:%d/%m/%Y} · parametrización y configuración hasta el {fin_param:%d/%m/%Y} · "
                f"paralelo con el aplicativo actual hasta el {CORTE:%d/%m/%Y} · salida en vivo {SALIDA:%d/%m/%Y}. "
                "Duración en días calendario, inclusive. Fechas estimadas: base para la salida en vivo.")
    ws["A2"].font = F_SUB
    write_legend(ws, 3, start_col=2)
    write_header(ws, 5, COLS_MOD)
    r = 6
    for t in tareas:
        vals = [t.id, t.fase_nombre, t.titulo, t.desc, t.resp, t.apoyo, t.ini, t.fin, f"=H{r}-G{r}+1",
                t.dep, t.crit, t.estado, t.obs]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = F_NORMAL
            c.alignment = A_WRAP
            c.border = BORDER
        ws.cell(row=r, column=2).fill = fill(FASES[t.fase][1])
        for j in (7, 8):
            ws.cell(row=r, column=j).number_format = DATE_FMT
            ws.cell(row=r, column=j).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r, column=9).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="top")
        if t.hito:
            for j in (1, 3):
                ws.cell(row=r, column=j).fill = fill(HITO_FILL)
                ws.cell(row=r, column=j).font = F_BOLD
        r += 1
    last = r - 1
    set_widths(ws, [6, 26, 44, 58, 34, 24, 12, 12, 10, 16, 48, 14, 42])
    ws.freeze_panes = "D6"
    ws.auto_filter.ref = f"A5:{get_column_letter(len(COLS_MOD))}{last}"
    add_estado_validation(ws, "L", 6, last)
    print_setup(ws, "5:5")
    return ws


# ---------------------------------------------------------------------------
# Hoja Consolidado (tabla + Gantt semanal + capacitaciones)
# ---------------------------------------------------------------------------
COLS_CONS = ["Nº", "Módulo", "Fase", "Tarea", "Responsable (cooperativa)", "Apoyo IngenIA",
             "Fecha inicio", "Fecha fin", "Duración", "Hito", "Estado"]


def todas_las_tareas():
    todas = [("Gestión del proyecto", t) for t in GESTION]
    for nombre, tareas, *_ in MODULOS:
        todas += [(nombre, t) for t in tareas]
    orden_mod = {"Gestión del proyecto": 0, "Nómina": 1, "Contabilidad": 2, "Comercial": 3, "Cartera financiera": 4}
    todas.sort(key=lambda x: (x[1].ini, x[1].fin, orden_mod[x[0]], x[1].id))
    return todas


def build_consolidado(wb):
    ws = wb.create_sheet("Consolidado")
    ws.sheet_properties.tabColor = HITO_STRONG
    todas = todas_las_tareas()

    ws["A1"] = f"Consolidado — Plan de implementación IngenIA365ERP · {COOP}"
    ws["A1"].font = F_TITLE
    ws["A2"] = (f"Todas las tareas ordenadas por fecha de inicio. Cronograma semanal: cada columna es la semana que "
                f"empieza el sábado indicado (sábado a viernes), del {SABADOS[0]:%d/%m} al {SABADOS[-1] + timedelta(days=6):%d/%m/%Y}. "
                f"El sombreado se calcula con formato condicional a partir de las fechas: si cambia una fecha, el cronograma se actualiza. "
                f"Corte del paralelo {CORTE:%d/%m} · salida en vivo {SALIDA:%d/%m/%Y}.")
    ws["A2"].font = F_SUB
    write_legend(ws, 3, start_col=2)

    n_fixed = len(COLS_CONS)
    first_week_col = n_fixed + 1
    # fila 5: etiquetas S1..S12 ; fila 6: encabezados + fecha del sábado
    for i, sab in enumerate(SABADOS):
        col = first_week_col + i
        c = ws.cell(row=5, column=col, value=f"S{i + 1}")
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = A_CENTER
        c.border = BORDER
    c = ws.cell(row=5, column=first_week_col - 1, value="Semana →")
    c.font = F_BOLD
    c.alignment = Alignment(horizontal="right", vertical="center")
    write_header(ws, 6, COLS_CONS)
    for i, sab in enumerate(SABADOS):
        col = first_week_col + i
        c = ws.cell(row=6, column=col, value=sab)
        c.number_format = "dd/mm"
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = A_CENTER
        c.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = 6.5

    r = 7
    for modulo, t in todas:
        vals = [t.id, modulo, t.fase_nombre, t.titulo, t.resp, t.apoyo, t.ini, t.fin, f"=H{r}-G{r}+1",
                "◆" if t.hito else "", t.estado]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = F_BOLD if (t.hito and j in (1, 4, 10)) else F_NORMAL
            c.alignment = A_WRAP
            c.border = BORDER
        ws.cell(row=r, column=3).fill = fill(FASES[t.fase][1])
        for j in (7, 8):
            ws.cell(row=r, column=j).number_format = DATE_FMT
            ws.cell(row=r, column=j).alignment = Alignment(horizontal="center", vertical="top")
        for j in (1, 9, 10):
            ws.cell(row=r, column=j).alignment = Alignment(horizontal="center", vertical="top")
        if t.hito:
            ws.cell(row=r, column=10).fill = fill(HITO_FILL)
        for i in range(len(SABADOS)):
            ws.cell(row=r, column=first_week_col + i).border = BORDER
        r += 1
    last = r - 1

    # Formato condicional del Gantt: hito primero (detiene), luego cada fase.
    wk0 = get_column_letter(first_week_col)
    wkN = get_column_letter(first_week_col + len(SABADOS) - 1)
    rango = f"{wk0}7:{wkN}{last}"
    overlap = f"$G7<={wk0}$6+6,$H7>={wk0}$6"
    ws.conditional_formatting.add(
        rango, FormulaRule(formula=[f'AND($J7="◆",{overlap})'], fill=fill(HITO_FILL), stopIfTrue=True))
    for key, (nombre, light, strong) in FASES.items():
        ws.conditional_formatting.add(
            rango, FormulaRule(formula=[f'AND($C7="{nombre}",{overlap})'], fill=fill(strong)))

    set_widths(ws, [6, 18, 26, 46, 32, 22, 12, 12, 9, 6, 14])
    ws.freeze_panes = "E7"
    ws.auto_filter.ref = f"A6:{get_column_letter(n_fixed)}{last}"
    add_estado_validation(ws, "K", 7, last)

    # Calendario de capacitaciones de los sábados
    r = last + 3
    ws.cell(row=r, column=1, value="Calendario de capacitaciones de los sábados").font = F_H2
    r += 1
    ws.cell(row=r, column=1, value=(
        "Sesiones por proceso, presenciales o virtuales; si dos módulos coinciden el mismo sábado, se hacen en paralelo "
        "o en bloques distintos. Horario y modalidad: acordado entre las partes. No se programan sesiones en festivos "
        "(12/10, 02/11, 16/11 y 08/12).")).font = F_SUB
    r += 1
    write_header(ws, r, ["Fecha", "Día", "Módulo(s)", "Sesiones (Nº y tema)", "Horario y modalidad", "Observaciones"])
    r += 1
    sesiones = {}
    for modulo, t in todas:
        if t.fase == "CAP":
            sesiones.setdefault(t.ini, []).append((modulo, t))
    for sab in SABADOS:
        items = sesiones.get(sab, [])
        mods = ", ".join(dict.fromkeys(m for m, _ in items)) or "—"
        temas = "\n".join(f"{t.id}: {t.tarea}" for _, t in items) or "Sin sesión programada (reservado para refuerzo o recuperación)"
        obs = "Corte del paralelo y comité go / no-go" if sab == CORTE else ""
        if sab == HOY:
            obs = "Hoy (19/09): acuerdo de horario y modalidad con gerencia"
        row = [sab, "sábado", mods, temas, "Acordado entre las partes", obs]
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = F_NORMAL
            c.alignment = A_WRAP
            c.border = BORDER
        ws.cell(row=r, column=1).number_format = DATE_FMT
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Sesiones entre semana (día por acordar)").font = F_H2
    r += 1
    write_header(ws, r, ["Fecha", "Día", "Módulo(s)", "Sesiones (Nº y tema)", "Horario y modalidad", "Observaciones"])
    r += 1
    for fecha in sorted(k for k in sesiones if k not in SABADOS):
        items = sesiones[fecha]
        row = [fecha, DIAS[fecha.weekday()], ", ".join(dict.fromkeys(m for m, _ in items)),
               "\n".join(f"{t.id}: {t.tarea}" for _, t in items), "Acordado entre las partes",
               "Fecha tentativa; se confirma en el seguimiento semanal"]
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = F_NORMAL
            c.alignment = A_WRAP
            c.border = BORDER
        ws.cell(row=r, column=1).number_format = DATE_FMT
        r += 1

    print_setup(ws, "5:6")
    return ws


# ---------------------------------------------------------------------------
# Hoja Resumen
# ---------------------------------------------------------------------------
def build_resumen(wb):
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_properties.tabColor = NAVY
    set_widths(ws, [30, 34, 34, 44, 34, 22, 18, 18])
    r = 1

    def title(text):
        nonlocal r
        ws.cell(row=r, column=1, value=text).font = F_TITLE
        r += 1

    def sub(text):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        c = ws.cell(row=r, column=1, value=text)
        c.font = F_SUB
        c.alignment = A_WRAP
        ws.row_dimensions[r].height = 32
        r += 1

    def h2(text):
        nonlocal r
        r += 1
        ws.cell(row=r, column=1, value=text).font = F_H2
        r += 1

    def par(text, cols=8, font=F_NORMAL, fill_=None, height=None):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=cols)
        c = ws.cell(row=r, column=1, value=text)
        c.font = font
        c.alignment = A_WRAP
        if fill_:
            c.fill = fill_
        ws.row_dimensions[r].height = height or max(18, 16 * (1 + len(text) // 150))
        r += 1

    def table(headers, rows, date_cols=()):
        nonlocal r
        write_header(ws, r, headers)
        r += 1
        for row in rows:
            for j, v in enumerate(row, start=1):
                c = ws.cell(row=r, column=j, value=v)
                c.font = F_NORMAL
                c.alignment = A_WRAP
                c.border = BORDER
                if j in date_cols:
                    c.number_format = DATE_FMT
                    c.alignment = Alignment(horizontal="center", vertical="top")
            r += 1

    title(f"Plan de implementación de IngenIA365ERP — {COOP}")
    sub(f"Versión 1 · {HOY:%d/%m/%Y} · elaborado por IngenIA 365 (Jorman Copete) para la Gerencia y el equipo de {COOP}. "
        "Se mantiene en este libro: los cambios acordados en el seguimiento semanal se registran aquí.")
    r += 1
    par("ADVERTENCIA: los tiempos de este plan son estimados; sin embargo, constituyen la base de trabajo para salir en vivo "
        f"el martes {SALIDA:%d/%m/%Y} únicamente con IngenIA365ERP. El paralelo con el aplicativo actual termina el sábado "
        f"{CORTE:%d/%m/%Y}. Cualquier desvío se registra en la bitácora y se escala a Gerencia en el seguimiento semanal.",
        font=F_WARN, fill_=FILL_WARN, height=48)

    h2("1. Propósito")
    par(f"Poner en operación IngenIA365ERP en {COOP} en los módulos de Nómina, Contabilidad, Comercial y Cartera financiera, "
        "con el personal capacitado y los datos parametrizados y conciliados, de modo que desde el 01/12/2026 toda la operación "
        "se registre únicamente en el nuevo sistema. El plan combina tres etapas por módulo: parametrización y configuración, "
        "capacitación por proceso y un paralelo con el aplicativo actual que permite interiorizar el manejo y detectar "
        "inconsistencias antes de la salida en vivo.", height=64)

    h2("2. Alcance por módulo")
    table(["Módulo", "Inicio", "Fin de parametrización", "Paralelo con el aplicativo actual", "Corte del paralelo", "Nº de tareas", "Hoja"],
          [[n, ini, fp, f"del {fp + timedelta(days=1):%d/%m} al {CORTE:%d/%m}", CORTE, len(tareas), n]
           for n, tareas, ini, fin, fp in MODULOS], date_cols=(2, 3, 5))
    par("Nómina: planes (quincenal), conceptos, parámetros legales 2026, retención por plan, EPS/ARL/pensiones/cesantías/cajas, "
        "empleados, períodos, novedades, liquidación (calcular → revisar → aprobar; la aprobación genera el comprobante contable NM) y reportes.", height=34)
    par("Contabilidad (entrega E1 en producción desde el 19/09): configuración inicial con el catálogo PUC Solidario (CUIF de la "
        "Supersolidaria, 2.110 cuentas) o PUC Comercial, nivel de movimiento 5 o 6, plan de cuentas con auxiliares propios, tipos de comprobante, "
        "períodos, comprobantes y borradores con control de cuatro ojos. Pendiente del proveedor: informes contables y estados financieros "
        "(entrega E2) e integración contable automática de Cartera, Inventario/Comercial, Tesorería y CDT (entrega E3), ambas con fecha por confirmar.", height=64)
    par("Comercial (menú Inventario y Tesorería): productos, grupos, bodegas, ubicaciones, puntos de venta, turnos, vendedores, descuentos, "
        "listas de precios, comisiones, tipos de movimiento, movimientos, facturación, kardex e inventario valorizado; cheques, facturas CxP/CxC, flujo de caja.", height=34)
    par("Cartera financiera (créditos y aportes): líneas de crédito, códigos de movimiento, tasas, provisión y calificación, zonas, SIPLA, "
        "periodicidad, cuentas de cartera, asociados, solicitudes, desembolsos, recaudos, mora y cobro, causación, calificación, descuento por "
        "nómina, extractos y cartera por edades. Fuera de alcance en esta implementación: ahorros a la vista y CDT (decisión del 19/09).", height=48)
    par("Transversal: usuarios, roles y permisos (administrador, Operador —digita, no contabiliza—, sólo lectura, Auditor), segundo factor "
        "obligatorio, sucursales, auditoría, maestros (personas, agencias, ciudades, bancos, centros de costo, cargos) y manual en línea.", height=34)

    h2("3. Fechas clave e hitos")
    hitos = [(m, t) for m, t in todas_las_tareas() if t.hito]
    table(["Fecha", "Día", "Hito", "Módulo", "Nº"],
          [[t.ini, DIAS[t.ini.weekday()], t.tarea, m, t.id] for m, t in hitos], date_cols=(1,))
    par(f"Festivos en el período (sin capacitaciones): lunes 12/10, lunes 02/11, lunes 16/11 y martes 08/12. "
        f"Hoy es sábado {HOY:%d/%m/%Y}: la primera semana de Nómina ya transcurrió y sus tareas figuran «Por confirmar».", height=34)

    h2("4. Roles y personas")
    table(["Rol", "Nombre", "Módulo / ámbito", "Responsabilidad en el plan"], [
        ["Gerencia", COOP, "Todos", "Patrocina el proyecto, designa a las personas, decide la salida en vivo (go / no-go)."],
        ["Gestor del proyecto (por la empresa)", "Por designar (G01)", "Todos",
         "Vela por el cumplimiento de las tareas y fechas; lleva la bitácora; dirige el seguimiento semanal; escala desvíos."],
        ["Contadora", "Rafaela Lastra España", "Contabilidad (líder); apoyo a Nómina, Comercial y Cartera",
         "Decide la configuración contable, valida el catálogo, define las cuentas por concepto y de cartera/inventario, concilia saldos al 30/11."],
        ["Auxiliar contable", "Por designar", "Contabilidad (2ª persona)", "Digita comprobantes y terceros; la contadora contabiliza (cuatro ojos)."],
        ["Equipo Nómina", "2 personas, por designar", "Nómina", "Parametriza, liquida en paralelo, compara y registra diferencias."],
        ["Equipo Comercial", "2 personas, por designar", "Comercial / Inventario / Tesorería", "Parametriza, registra la muestra del paralelo, toma física e inventario inicial."],
        ["Equipo Cartera", "2 personas, por designar", "Cartera financiera", "Parametriza, registra la muestra del paralelo, precarga créditos y aportes."],
        ["IngenIA 365", "Jorman Copete", "Proveedor", "Capacita, acompaña la parametrización, entrega E2 y E3, soporta la salida en vivo."],
    ])
    par("Regla fijada por el dueño del producto: mínimo 2 personas por módulo en capacitación e implementación (continuidad y cumplimiento "
        "de fechas) y una persona de la empresa encargada de la gestión del proyecto.", height=34)

    h2("5. Capacitación")
    par("Los sábados, por proceso, presencial o virtual; también otros días de la semana según acuerdo. Horario y modalidad: acordado entre las "
        "partes (gerencia e IngenIA, 19/09). El calendario completo de sesiones está en la hoja Consolidado. Cada sesión se hace con los datos "
        "reales de la cooperativa y termina con un ejercicio en el sistema; el Manual en línea (/manual) queda como referencia permanente.", height=48)

    h2("6. Regla del paralelo")
    par("Nómina: cada quincena de octubre y noviembre se liquida en ambos sistemas y se compara empleado por empleado (ciclo de calibración con la "
        "2ª quincena de septiembre). Contabilidad, Comercial y Cartera: no se replica toda la operación, sólo las transacciones más importantes y las "
        "que cubren todos los casos de la entidad. En cada ciclo hay una tarea de «comparar con el aplicativo actual y registrar diferencias»; toda "
        "diferencia va a la bitácora con causa, responsable y cierre. El 28/11 se corta el paralelo y el comité go / no-go decide la salida en vivo.", height=64)

    h2("7. Riesgos y mitigación")
    table(["Riesgo", "Impacto", "Mitigación", "Responsable"], [
        ["Disponibilidad de las 2 personas por módulo (vacaciones, cargas de cierre, retiros)", "Retraso de la parametrización o del paralelo",
         "Suplentes designados desde G01; sesiones virtuales grabadas cuando la modalidad lo permita; seguimiento semanal con escalamiento a Gerencia", "Gerencia / Gestor"],
        ["Dependencia del proveedor: entregas E2 (informes contables) y E3 (integración automática) sin fecha",
         "Comparación de saldos más lenta; contabilización manual de Cartera y Comercial",
         "Confirmar fechas en la primera semana (G07); comparar por comprobantes y exportación; comprobante resumen manual (C35) hasta E3", "IngenIA / Contadora"],
        ["Calidad de los datos del aplicativo actual (terceros duplicados, cuentas sin homologar, productos inactivos)", "Diferencias falsas en el paralelo; retrabajo",
         "Levantamiento y depuración en la fase de preparación; mapa de cuentas y de conceptos; verificación cruzada antes de cada hito de parametrización", "Equipos / Gestor"],
        ["Carga de saldos iniciales concentrada entre el 28/11 y el 04/12", "Salida en vivo con saldos incompletos",
         "Precarga de créditos y aportes con saldos al 27/11 y ajuste de los movimientos del 28 al 30/11; si el volumen supera lo supuesto, importación asistida por IngenIA", "Equipo Cartera / Contadora"],
        ["Festivos (12/10, 02/11, 16/11, 08/12) y eventos de la cooperativa en sábados", "Menos días efectivos; sesiones perdidas",
         "No se programan sesiones en festivos; los ciclos ya descuentan esos días; el sábado sin sesión asignada se usa para recuperar", "Gestor"],
        ["Cierre de noviembre del aplicativo actual después del 01/12", "Saldos iniciales provisionales",
         "Comprobante de apertura con saldos preliminares y comprobante de ajuste posterior (C37)", "Contadora"],
        ["Alcance por confirmar: información para PILA (N35) y facturación electrónica DIAN (M27)", "Procesos que seguirían fuera del sistema",
         "Confirmar con IngenIA en octubre y documentar el procedimiento transitorio", "IngenIA / Equipos"],
        ["Segundo factor obligatorio: personas sin celular compatible", "Usuarios sin acceso",
         "Definir método (TOTP o passkey) al crear cada usuario (G03)", "Administrador"],
        ["Resistencia al cambio de hábitos", "Uso paralelo informal del aplicativo viejo después del 01/12",
         "Comunicados de inicio y de salida en vivo; aplicativo actual en solo lectura desde el 01/12; acompañamiento la primera semana", "Gerencia / Gestor"],
    ])

    h2("8. Supuestos (marcados «supuesto» hasta que se confirmen)")
    table(["Supuesto", "Origen / estado"], [
        ["Nómina quincenal, con un solo plan de nómina; primera quincena en paralelo formal = 1 al 15 de octubre", "Decisión del dueño del producto (19/09)"],
        ["El paralelo contable incluye las nóminas de octubre y noviembre y cubre ventas, compras, caja, bancos, causaciones y egresos", "Decisión del dueño del producto (19/09)"],
        ["Cartera: sólo créditos y aportes; ahorros a la vista y CDT fuera de alcance", "Decisión del dueño del producto (19/09)"],
        ["Volumen: menos de 500 asociados y menos de 300 créditos vigentes; número de empleados y de productos por confirmar", "Supuesto — confirmar en el levantamiento (M02, K02)"],
        ["Saldos iniciales al 30/11 digitados desde el aplicativo actual y conciliados por la contadora y el responsable de cada módulo", "Decisión del dueño del producto (19/09)"],
        ["Horario y modalidad de las capacitaciones", "Acordado entre las partes (gerencia e IngenIA, 19/09)"],
        ["Gestor del proyecto, auxiliar contable y equipos por módulo", "Por designar (G01)"],
        ["Contabilidad con catálogo PUC Solidario (CUIF); nivel de movimiento por decidir con la contadora (C02)", "Supuesto — decisión en C02"],
        ["Entrega E1 de Contabilidad en producción desde el 19/09; E2 y E3 con fecha por confirmar", "Estado del producto al 19/09"],
        ["Los movimientos de Cartera y Comercial se contabilizan con comprobantes resumen hasta la entrega E3", "Supuesto — procedimiento en C35"],
    ])

    h2("9. Cómo usar este libro")
    par("Cada módulo tiene su hoja con tareas por fase (Nº, fase, tarea, descripción, responsable, apoyo, fechas, duración, dependencia, criterio de listo, "
        "estado y observaciones). El estado se elige de una lista (Pendiente, En curso, Terminada, Bloqueada, Por confirmar) y se colorea solo. "
        "La hoja Consolidado reúne todas las tareas ordenadas por fecha con el cronograma semanal: el sombreado sale del formato condicional, así que al "
        "cambiar una fecha el cronograma se actualiza. Los filtros están activos en los encabezados y los paneles quedan inmovilizados. Sin macros.", height=64)
    r += 1
    write_legend(ws, r)
    ws.freeze_panes = "A4"
    print_setup(ws, None)
    return ws


# ---------------------------------------------------------------------------
# Generación y verificación
# ---------------------------------------------------------------------------
def generar(ruta):
    wb = Workbook()
    build_resumen(wb)
    for nombre, tareas, ini, fin, fp in MODULOS:
        build_modulo(wb, nombre, tareas, ini, fin, fp)
    build_consolidado(wb)
    wb.properties.title = f"Plan de implementación IngenIA365ERP — {COOP}"
    wb.properties.creator = "IngenIA 365"
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta


def verificar(ruta):
    wb = load_workbook(ruta)
    print(f"Archivo: {ruta}")
    print("Hojas:", ", ".join(wb.sheetnames))
    for nombre, tareas, ini, fin, fp in MODULOS + [("Gestión del proyecto", GESTION, D(9, 19), D(12, 4), None)]:
        por_fase = {}
        for t in tareas:
            por_fase[t.fase_nombre] = por_fase.get(t.fase_nombre, 0) + 1
        hitos = sum(1 for t in tareas if t.hito)
        print(f"\n{nombre}: {len(tareas)} tareas ({hitos} hitos) · {min(t.ini for t in tareas):%d/%m} → {max(t.fin for t in tareas):%d/%m}"
              + (f" · fin parametrización {fp:%d/%m}" if fp else ""))
        for f, n in por_fase.items():
            print(f"   {f}: {n}")
        for t in tareas:
            assert ini <= t.ini and t.fin <= D(12, 4), f"{t.id} fuera de la ventana del módulo"
        ids = [t.id for t in tareas]
        assert len(ids) == len(set(ids)), f"{nombre}: Nº repetido"
    ws = wb["Consolidado"]
    print(f"\nConsolidado: {ws.max_row} filas · reglas de formato condicional: {len(ws.conditional_formatting)}")
    caps = [t for _, ts, *_ in MODULOS for t in ts if t.fase == "CAP"]
    print(f"Sesiones de capacitación: {len(caps)} ({sum(1 for t in caps if t.ini in SABADOS)} en sábado, "
          f"{sum(1 for t in caps if t.ini not in SABADOS)} entre semana)")
    fest = [t for _, ts, *_ in MODULOS for t in ts if t.fase == 'CAP' and t.ini in FESTIVOS]
    assert not fest, "Hay capacitaciones en festivo: " + ", ".join(t.id for t in fest)


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    destino = Path(args[0]) if args else Path(__file__).with_name("plan-de-implementacion-2026.xlsx")
    generar(destino)
    verificar(destino)
    if "--descargas" in sys.argv:
        descargas = Path.home() / "Downloads" / destino.name
        shutil.copyfile(destino, descargas)
        print(f"Copia: {descargas}")
